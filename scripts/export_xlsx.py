"""Write an internship sweep to an Excel workbook.

Usage: python export_xlsx.py <sweep.json> <output.xlsx>

sweep.json:
  {
    "run_date": "2026-10-21",
    "params": {"Role": "finance intern", "Region": "Vancouver, BC"},
    "notes": ["Indeed stopped at page 2: CAPTCHA"],
    "rows": [{"role": "...", "company": "...", "location": "...", "work_mode": "...",
              "term": "...", "posted": "2026-10-15", "source": "LinkedIn",
              "link": "https://...", "tracked": "", "notes": ""}]
  }

"posted" in YYYY-MM-DD becomes an Excel date; any other text stays text.
Only http and https links become hyperlinks.
Never overwrites: if the output name is taken, " (2)", " (3)" and so on is added.
Reopens the saved file and exits with an error if the row count is wrong.
Prints the path written and the row count.
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("#", None, 5),
    ("Role", "role", 38),
    ("Company", "company", 26),
    ("Location", "location", 24),
    ("Work mode", "work_mode", 14),
    ("Term", "term", 16),
    ("Posted", "posted", 13),
    ("Source", "source", 16),
    ("Link", "link", 10),
    ("Already tracked", "tracked", 22),
    ("Notes", "notes", 40),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")


def unique_path(path):
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.stem} ({n}){path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def set_text(cell, value):
    cell.value = None if value == "" else value
    # openpyxl reads a leading = as a formula; keep scraped text as text.
    if isinstance(value, str) and value and value[0] in "=+-@":
        cell.data_type = "s"


def as_date(value):
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def build_results(wb, rows):
    ws = wb.active
    ws.title = "Results"
    headers = [name for name, _, _ in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    col_of = {key: i for i, (_, key, _) in enumerate(COLUMNS, start=1) if key}
    for i, row in enumerate(rows, start=1):
        r = i + 1
        ws.cell(row=r, column=1, value=i)
        for key, col in col_of.items():
            cell = ws.cell(row=r, column=col)
            value = row.get(key, "")
            if key == "link":
                if isinstance(value, str) and value.lower().startswith(("http://", "https://")):
                    cell.value = "Open"
                    cell.hyperlink = value
                    cell.font = LINK_FONT
                elif value:
                    set_text(cell, value)
            elif key == "posted" and as_date(value):
                cell.value = as_date(value)
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="left")
            else:
                set_text(cell, value)
        ws.cell(row=r, column=col_of["notes"]).alignment = Alignment(wrap_text=True, vertical="top")

    for col, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_search(wb, data):
    ws = wb.create_sheet("Search")
    ws.append(["Run date", data["run_date"]])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Parameter", "Value"])
    for cell in ws[3]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for key, value in data.get("params", {}).items():
        r = ws.max_row + 1
        set_text(ws.cell(row=r, column=1), key)
        set_text(ws.cell(row=r, column=2), value)
    if data.get("notes"):
        ws.append([])
        ws.append(["Run notes"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for note in data["notes"]:
            set_text(ws.cell(row=ws.max_row + 1, column=1), note)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 70


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    data = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out = unique_path(Path(sys.argv[2]))
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    build_results(wb, data["rows"])
    build_search(wb, data)
    wb.save(out)

    saved = load_workbook(out)["Results"]
    got = saved.max_row - 1
    if got != len(data["rows"]):
        sys.exit(f"Row count mismatch in {out}: wrote {got}, expected {len(data['rows'])}")
    print(out)
    print(f"{got} rows, read back and matched")


if __name__ == "__main__":
    main()
